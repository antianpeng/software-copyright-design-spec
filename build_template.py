# -*- coding: utf-8 -*-
"""
软著申请材料一键生成脚本（通用模板）
生成：申请表(.doc) + 合并源代码(.docx) + 软件设计说明书(.docx)

用法：
    1. 修改下方 CONFIG 区域的所有 TODO 项
    2. 修改 generate_design_spec() 中的项目内容
    3. python build_template.py

模板说明：
    - 所有 TODO 标记的地方都需要根据你的项目填写
    - 工具函数和文档结构是通用的，不需要改
    - 设计说明书的章节结构是标准的，只需填充内容
"""
import os
import sys
import subprocess

sys.stdout.reconfigure(encoding='utf-8')

from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# ═══════════════════════════════════════════════════════════════════════
# CONFIG — 修改这里（所有 TODO 项必须填写）
# ═══════════════════════════════════════════════════════════════════════

# TODO: 软件全称，必须以"软件"、"系统"或"平台"结尾
SOFTWARE_NAME = "XXX管理系统软件"
# TODO: 版本号，格式 VX.X 或 VX.X.X
VERSION = "V1.0"
# TODO: 项目代码根目录
PROJECT_ROOT = r"D:\code\your-project"
# TODO: 申请表模板 .doc 文件路径（从模板目录复制一份）
TEMPLATE_DOC = r"D:\2026软著\模板\模板文件名.doc"
# TODO: 输出目录
OUTPUT_DIR = r"D:\2026软著\你的软件名称"

# 工具路径（一般不需要改）
MMDC = "mmdc"  # 或完整路径: r"C:\Users\{user}\AppData\Roaming\npm\mmdc.cmd"
DOT = "dot"    # 或完整路径: r"C:\Program Files\Graphviz\bin\dot.exe"
DIAGRAMS_DIR = os.path.join(OUTPUT_DIR, "diagrams")

# ═══════════════════════════════════════════════════════════════════════
# 申请表字段值（按模板26行顺序填写）
# ═══════════════════════════════════════════════════════════════════════

# TODO: 填写26个字段的值，空字符串表示留空让用户手动填
# 前4项（著作权人等）通常留空
FILL_VALUES = {
    0: "",                    # TODO: 著作权人姓名或名称
    1: "",                    # TODO: 公司成立日期
    2: "",                    # TODO: 营业执照注册号
    3: "",                    # TODO: 公司地址
    4: SOFTWARE_NAME,         # 软件名称（自动填入）
    5: VERSION,               # 版本号（自动填入）
    6: "否",                  # 是否升级版本
    7: "",                    # 软件简称（选填）
    8: "应用软件",            # 软件分类
    9: "",                    # TODO: 开发完成日期（格式 YYYY-MM-DD，需在公司成立后3个月以上）
    10: "未发表",             # 发表状态
    11: "",                   # 首次发表时间（未发表则空）
    12: "城市：",             # 首次发表地点
    13: "",                   # TODO: 开发硬件环境（≤50字，如"Intel CPU 8核、内存16GB、硬盘256GB"）
    14: "",                   # TODO: 运行硬件环境（≤50字，如"Intel CPU 2核、内存4GB、硬盘50GB"）
    15: "",                   # TODO: 开发操作系统（如"Windows 10"）
    16: "",                   # TODO: 开发环境/工具（如"开发环境: Windows 10/开发工具: IntelliJ IDEA"）
    17: "",                   # TODO: 运行平台/操作系统（如"Windows 10/11"）
    18: "",                   # TODO: 运行支撑环境（如"JDK 11、MySQL 5.7"）
    19: "",                   # TODO: 编程语言（如"Java, JavaScript"）
    20: "",                   # TODO: 源程序量（总行数，用 wc -l 统计）
    21: "",                   # TODO: 开发目的（50字以内）
    22: "",                   # TODO: 面向领域/行业
    23: "",                   # TODO: 主要功能（500-1300字，详细介绍每个模块功能）
    24: "",                   # TODO: 技术特点（≤100字）
    25: "应用软件",           # 技术特点标签
}

# ═══════════════════════════════════════════════════════════════════════
# 源代码文件列表（相对于 PROJECT_ROOT）
# ═══════════════════════════════════════════════════════════════════════

# TODO: 从项目中选取核心源代码文件，总行数需 ≥3000行（60页×50行）
# 选取原则：核心业务文件（Controller/Service/Model）+ 配置文件 + 前端路由
SELECTED_FILES = [
    # === 配置文件 ===
    "pom.xml",                          # TODO: 改为你的项目配置文件
    # "package.json",                    # 前端项目
    # "go.mod",                         # Go 项目
    # "requirements.txt",               # Python 项目

    # === 后端核心文件 ===
    # TODO: 列出你的 Controller 文件
    # "src/main/java/com/example/controller/UserController.java",
    # TODO: 列出你的 Service 文件
    # "src/main/java/com/example/service/UserService.java",
    # TODO: 列出你的 Entity/Model 文件
    # "src/main/java/com/example/entity/User.java",

    # === 前端核心文件（如有）===
    # "src/router/index.js",
    # "src/views/Login.vue",
    # "src/api/user.js",
]

# ═══════════════════════════════════════════════════════════════════════
# 工具函数（通用，不需要修改）
# ═══════════════════════════════════════════════════════════════════════

def set_run_font(run, name="宋体", size=9, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")}/>')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), name)


def set_cell_shading(cell, color):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def set_page_header(section, text):
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    p.clear()
    run = p.add_run(text)
    set_run_font(run, "宋体", 9)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_heading_styled(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "黑体"
        r = run._element
        rPr = r.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = parse_xml(f'<w:rFonts {nsdecls("w")}/>')
            rPr.insert(0, rFonts)
        rFonts.set(qn('w:eastAsia'), "黑体")
    return h


def add_para(doc, text, font_name="宋体", font_size=12, bold=False, indent=0, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, font_name, font_size, bold)
    if indent:
        p.paragraph_format.first_line_indent = Cm(indent)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_table_from_rows(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.paragraphs[0].clear()
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, "宋体", 10, bold=True)
        set_cell_shading(cell, "D9E2F3")
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.paragraphs[0].clear()
            run = cell.paragraphs[0].add_run(str(val))
            set_run_font(run, "宋体", 10)
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)
    return table


# ═══════════════════════════════════════════════════════════════════════
# 图表生成（通用引擎 + 项目定制内容）
# ═══════════════════════════════════════════════════════════════════════

def run_mermaid(name, mmd_content):
    os.makedirs(DIAGRAMS_DIR, exist_ok=True)
    mmd_path = os.path.join(DIAGRAMS_DIR, f"{name}.mmd")
    png_path = os.path.join(DIAGRAMS_DIR, f"{name}.png")
    with open(mmd_path, "w", encoding="utf-8") as f:
        f.write(mmd_content)
    result = subprocess.run(
        [MMDC, "-i", mmd_path, "-o", png_path, "-b", "white", "-w", "1200"],
        capture_output=True, text=True, timeout=60
    )
    if result.returncode == 0 and os.path.exists(png_path):
        print(f"  [OK] {name}.png")
    else:
        print(f"  [FAIL] {name}: {result.stderr[:200]}")
    return png_path


def run_graphviz(name, dot_content):
    os.makedirs(DIAGRAMS_DIR, exist_ok=True)
    dot_path = os.path.join(DIAGRAMS_DIR, f"{name}.dot")
    png_path = os.path.join(DIAGRAMS_DIR, f"{name}.png")
    with open(dot_path, "w", encoding="utf-8") as f:
        f.write(dot_content)
    result = subprocess.run(
        [DOT, "-Tpng", "-Gdpi=150", dot_path, f"-o{png_path}"],
        capture_output=True, text=True, timeout=30
    )
    if result.returncode == 0 and os.path.exists(png_path):
        print(f"  [OK] {name}.png")
    else:
        print(f"  [FAIL] {name}: {result.stderr[:200]}")
    return png_path


def generate_diagrams():
    """
    生成全部10张图表。

    TODO: 根据你的项目修改以下图表内容：
    - 01 系统架构图：修改层次和节点
    - 02 模块分解图：修改模块名称
    - 03-07 流程图：根据实际业务逻辑修改
    - 08 时序图：修改参与者和消息
    - 09 ER图：修改实体和关系
    - 10 数据流图：修改数据流向

    注意：Graphviz 的 subgraph label 必须用英文，否则中文乱码！
    """
    print("\n=== 生成图表 ===")

    # ── 01 系统架构图 (Graphviz) ──
    # TODO: 根据你的系统架构修改层次和节点
    dot_arch = f'''digraph architecture {{
    rankdir=TB;
    node [shape=box, style="rounded,filled", fillcolor="#E8F4FD", fontname="Microsoft YaHei", fontsize=11];
    edge [fontname="Microsoft YaHei", fontsize=9];

    subgraph cluster_client {{ label="Client Layer"; style="rounded,dashed"; fillcolor="#F5F5F5";
        browser [label="浏览器"];
    }}
    subgraph cluster_server {{ label="Server Layer"; style="rounded,dashed"; fillcolor="#E8F5E9";
        backend [label="应用服务器"];
    }}
    subgraph cluster_data {{ label="Data Layer"; style="rounded,dashed"; fillcolor="#FCE4EC";
        db [label="数据库"];
    }}

    browser -> backend -> db;
}}'''
    run_graphviz("01_system_architecture", dot_arch)

    # ── 02 模块分解图 (Graphviz) ──
    # TODO: 根据你的项目模块修改
    dot_mod = '''digraph modules {
    rankdir=TB;
    node [shape=box, style="rounded,filled", fontname="Microsoft YaHei", fontsize=10];

    subgraph cluster_mod1 { label="Module Group 1"; style="rounded"; fillcolor="#E3F2FD"; color="#1565C0";
        mod_a [label="Module A" fillcolor="#BBDEFB"];
        mod_b [label="Module B" fillcolor="#BBDEFB"];
    }
    subgraph cluster_mod2 { label="Module Group 2"; style="rounded"; fillcolor="#E8F5E9"; color="#2E7D32";
        mod_c [label="Module C" fillcolor="#C8E6C9"];
        mod_d [label="Module D" fillcolor="#A5D6A7"];
    }

    mod_a -> mod_c;
    mod_b -> mod_d;
}'''
    run_graphviz("02_module_decomposition", dot_mod)

    # ── 03 登录流程 (Mermaid) ──
    # TODO: 根据你的登录逻辑修改
    run_mermaid("03_login_flow", """flowchart TD
    A[用户打开浏览器] --> B[显示登录页面]
    B --> C[输入账号密码]
    C --> D{验证}
    D -->|失败| E[提示错误]
    E --> C
    D -->|成功| F[生成Token]
    F --> G[跳转首页]""")

    # ── 04-07 核心业务流程 (Mermaid) ──
    # TODO: 根据你的核心业务模块修改，每个模块一个流程图
    # 建议选择最能体现业务逻辑的4-5个模块

    # 示例：模块1的CRUD流程
    run_mermaid("04_module1_flow", """flowchart TD
    A[进入模块1] --> B[查看列表]
    B --> C{选择操作}
    C -->|新增| D[填写信息]
    D --> E[校验]
    E -->|通过| F[保存]
    C -->|编辑| G[修改信息]
    G --> F
    C -->|删除| H[确认删除]
    H --> I[删除记录]""")

    # 示例：模块2的业务流程
    run_mermaid("05_module2_flow", """flowchart TD
    A[进入模块2] --> B[加载数据]
    B --> C{选择操作}
    C -->|操作1| D[执行处理]
    D --> E[更新状态]
    C -->|操作2| F[查询分析]
    F --> G[展示结果]""")

    # 示例：模块3的业务流程
    run_mermaid("06_module3_flow", """flowchart TD
    A[进入模块3] --> B[选择对象]
    B --> C[配置参数]
    C --> D{确认执行}
    D -->|是| E[执行操作]
    E --> F[返回结果]
    D -->|否| G[取消]""")

    # 示例：模块4的业务流程
    run_mermaid("07_module4_flow", """flowchart TD
    A[进入模块4] --> B[查看状态]
    B --> C{需要操作?}
    C -->|是| D[执行操作]
    D --> E[更新状态]
    C -->|否| F[返回列表]""")

    # ── 08 通信时序图 (Mermaid) ──
    # TODO: 根据你的系统通信方式修改参与者和消息
    run_mermaid("08_sequence", """sequenceDiagram
    participant C as 客户端
    participant S as 服务器
    participant D as 数据库

    C->>S: 请求
    S->>D: 查询数据
    D->>S: 返回结果
    S->>C: 响应数据""")

    # ── 09 数据库ER图 (Mermaid) ──
    # TODO: 根据你的数据库表结构修改实体和关系
    run_mermaid("09_database_er", """erDiagram
    USER {
        bigint id PK
        varchar username UK
        varchar password
        varchar email
        int status
    }
    ORDER {
        bigint id PK
        bigint user_id FK
        decimal amount
        int status
        datetime create_time
    }
    USER ||--o{ ORDER : "creates" """)

    # ── 10 数据流图 (Mermaid) ──
    # TODO: 根据你的系统数据流修改
    run_mermaid("10_data_flow", """flowchart LR
    subgraph Input
        A1[用户请求] --> B1[API接口]
    end
    subgraph Process
        B1 --> C1[业务处理]
        C1 --> C2[数据访问]
    end
    subgraph Storage
        C2 --> D1[(数据库)]
    end
    subgraph Output
        D1 --> E1[页面渲染]
    end""")

    print("图表生成完毕")


# ═══════════════════════════════════════════════════════════════════════
# 文档1：申请表 (.doc)
# ═══════════════════════════════════════════════════════════════════════

def generate_application_form():
    """复制模板 .doc 并替换字段值"""
    print("\n=== 生成申请表 ===")
    output_path = os.path.join(OUTPUT_DIR, f"{SOFTWARE_NAME}.doc")

    if os.path.exists(TEMPLATE_DOC):
        import shutil
        shutil.copy2(TEMPLATE_DOC, output_path)
        doc = Document(output_path)
        table = doc.tables[0]
        for idx, val in FILL_VALUES.items():
            if val == "":
                continue
            if idx < len(table.rows):
                cell = table.rows[idx].cells[1]
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.text = ""
                if cell.paragraphs:
                    cell.paragraphs[0].clear()
                    run = cell.paragraphs[0].add_run(val)
                    set_run_font(run, "微软雅黑", 14)
        doc.save(output_path)
        print(f"  [OK] {output_path} (从模板生成)")
    else:
        print(f"  [ERROR] 模板不存在: {TEMPLATE_DOC}")
        print(f"  请先复制一份模板 .doc 文件到该路径")


# ═══════════════════════════════════════════════════════════════════════
# 文档2：合并源代码 (.docx)
# ═══════════════════════════════════════════════════════════════════════

def generate_merged_code():
    """合并源代码为60页文档（前30+后30），每50行一页"""
    print("\n=== 生成合并源代码 ===")
    output_path = os.path.join(OUTPUT_DIR, f"{SOFTWARE_NAME}--源代码.docx")

    all_lines = []
    for rel_path in SELECTED_FILES:
        full_path = os.path.join(PROJECT_ROOT, rel_path)
        if not os.path.exists(full_path):
            print(f"  [SKIP] {rel_path}")
            continue
        try:
            with open(full_path, "r", encoding="utf-8", errors="ignore") as f:
                lines = [l.rstrip() for l in f if l.strip()]
            all_lines.append(f"// File: {rel_path}")
            all_lines.extend(lines)
            all_lines.append("")
        except Exception as e:
            print(f"  [SKIP] {rel_path}: {e}")

    if not all_lines:
        print("  [ERROR] 没有读取到任何源代码文件")
        return

    PAGE_SIZE = 50
    pages = [all_lines[i:i + PAGE_SIZE] for i in range(0, len(all_lines), PAGE_SIZE)]
    if len(pages) <= 60:
        selected_pages = pages
    else:
        selected_pages = pages[:30] + pages[-30:]

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    set_page_header(section, f"{SOFTWARE_NAME} {VERSION}")

    for page_idx, page_lines in enumerate(selected_pages):
        if page_idx > 0:
            doc.add_page_break()
        for line in page_lines:
            p = doc.add_paragraph()
            run = p.add_run(line)
            set_run_font(run, "Courier New", 8)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = Pt(10)

    doc.save(output_path)
    print(f"  [OK] {output_path} ({len(selected_pages)} pages, {len(all_lines)} lines)")


# ═══════════════════════════════════════════════════════════════════════
# 文档3：软件设计说明书 (.docx)
# ═══════════════════════════════════════════════════════════════════════

def _add_image(doc, diagram_name):
    """嵌入图表"""
    img_path = os.path.join(DIAGRAMS_DIR, f"{diagram_name}.png")
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(img_path, width=Inches(5.5))
    else:
        add_para(doc, f"[图表缺失: {diagram_name}.png]", "宋体", 10)


def generate_design_spec():
    """
    生成软件设计说明书。

    TODO: 根据你的项目修改以下内容：
    - 引言：项目背景、术语定义
    - 总体设计：技术栈表格、架构描述
    - 接口设计：API 接口表格
    - 模块设计：每个模块的功能说明、流程图、函数表、算法说明
    - 运行设计：运行流程、时序图、数据流图
    - 出错处理：异常处理策略
    - 数据库设计：ER 图、表结构
    """
    print("\n=== 生成软件设计说明书 ===")
    output_path = os.path.join(OUTPUT_DIR, f"{SOFTWARE_NAME}--软件设计说明书.docx")

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    # ── 封面 ──
    for _ in range(6):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(SOFTWARE_NAME)
    set_run_font(run, "黑体", 22, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("软件设计说明书")
    set_run_font(run, "黑体", 18)
    ver = doc.add_paragraph()
    ver.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = ver.add_run(f"版本：{VERSION}")
    set_run_font(run, "宋体", 14)
    doc.add_page_break()

    # ── 目录 ──
    add_heading_styled(doc, "目录", 1)
    for item in ["一、引言", "二、总体设计", "三、接口设计", "四、模块设计",
                  "五、运行设计", "六、出错处理设计", "七、数据库设计"]:
        add_para(doc, item, "宋体", 12)
    doc.add_page_break()

    # ═════════════════════════════════════════════════════════════════════
    # 一、引言
    # ═════════════════════════════════════════════════════════════════════
    add_heading_styled(doc, "一、引言", 1)

    add_heading_styled(doc, "1.1 编写目的", 2)
    # TODO: 填写编写目的
    add_para(doc, f"本文档是《{SOFTWARE_NAME}》的软件设计说明书，旨在详细描述系统的总体设计、模块划分、接口设计、算法逻辑和运行方案，为软件著作权登记提供技术文档支撑。", "宋体", 12, indent=0.74)

    add_heading_styled(doc, "1.2 项目背景", 2)
    # TODO: 填写项目背景（这个软件解决什么问题）
    add_para(doc, "TODO: 描述项目背景，这个软件解决什么问题，面向什么行业/场景。", "宋体", 12, indent=0.74)

    add_heading_styled(doc, "1.3 术语定义", 2)
    # TODO: 填写项目特有术语
    add_para(doc, "TODO: 列出项目特有的技术术语及其定义。", "宋体", 12, indent=0.74)

    doc.add_page_break()

    # ═════════════════════════════════════════════════════════════════════
    # 二、总体设计
    # ═════════════════════════════════════════════════════════════════════
    add_heading_styled(doc, "二、总体设计", 1)

    add_heading_styled(doc, "2.1 系统架构", 2)
    # TODO: 描述系统架构
    add_para(doc, "TODO: 描述系统的技术架构（B/S、C/S、微服务等）。", "宋体", 12, indent=0.74)
    _add_image(doc, "01_system_architecture")

    add_heading_styled(doc, "2.2 技术栈", 2)
    # TODO: 填写技术栈表格
    add_table_from_rows(doc, ["层次", "技术选型"], [
        ("后端框架", "TODO"),
        ("前端框架", "TODO"),
        ("数据库", "TODO"),
        ("缓存", "TODO"),
        ("其他", "TODO"),
    ], [4, 10])

    add_heading_styled(doc, "2.3 模块划分", 2)
    # TODO: 描述模块划分
    add_para(doc, "TODO: 描述系统的模块划分。", "宋体", 12, indent=0.74)
    _add_image(doc, "02_module_decomposition")

    doc.add_page_break()

    # ═════════════════════════════════════════════════════════════════════
    # 三、接口设计
    # ═════════════════════════════════════════════════════════════════════
    add_heading_styled(doc, "三、接口设计", 1)

    add_heading_styled(doc, "3.1 REST API接口", 2)
    # TODO: 填写 API 接口表格
    add_para(doc, "TODO: 描述系统的 REST API 接口设计。", "宋体", 12, indent=0.74)
    add_table_from_rows(doc, ["接口路径", "模块", "功能"], [
        ("/api/xxx/**", "TODO模块名", "TODO功能描述"),
    ], [4, 3, 6])

    add_heading_styled(doc, "3.2 其他接口", 2)
    # TODO: 填写 WebSocket / TCP / MQ 等其他接口
    add_para(doc, "TODO: 描述其他类型的接口（WebSocket、TCP、消息队列等）。", "宋体", 12, indent=0.74)

    doc.add_page_break()

    # ═════════════════════════════════════════════════════════════════════
    # 四、模块设计
    # ═════════════════════════════════════════════════════════════════════
    add_heading_styled(doc, "四、模块设计", 1)

    # TODO: 为每个核心业务模块添加以下4个小节
    # 建议选3-5个核心模块，每个模块包含：
    #   4.X.1 功能说明
    #   4.X.2 处理流程图
    #   4.X.3 关键函数表
    #   4.X.4 算法说明

    # 示例模块1
    add_heading_styled(doc, "4.1 模块1名称", 2)
    add_heading_styled(doc, "4.1.1 功能说明", 3)
    add_para(doc, "TODO: 描述这个模块的功能、管理的实体、支持的操作。", "宋体", 12, indent=0.74)
    add_heading_styled(doc, "4.1.2 处理流程图", 3)
    _add_image(doc, "04_module1_flow")
    add_heading_styled(doc, "4.1.3 关键函数", 3)
    add_table_from_rows(doc, ["函数名", "所属类", "功能说明"], [
        ("TODO()", "TODO_Controller", "TODO功能描述"),
    ], [3, 4, 7])
    add_heading_styled(doc, "4.1.4 算法说明", 3)
    add_para(doc, "TODO: 描述核心算法逻辑。", "宋体", 12, indent=0.74)

    # 示例模块2
    add_heading_styled(doc, "4.2 模块2名称", 2)
    add_heading_styled(doc, "4.2.1 功能说明", 3)
    add_para(doc, "TODO: 描述这个模块的功能。", "宋体", 12, indent=0.74)
    add_heading_styled(doc, "4.2.2 处理流程图", 3)
    _add_image(doc, "05_module2_flow")
    add_heading_styled(doc, "4.2.3 关键函数", 3)
    add_table_from_rows(doc, ["函数名", "所属类", "功能说明"], [
        ("TODO()", "TODO_Service", "TODO功能描述"),
    ], [3, 4, 7])
    add_heading_styled(doc, "4.2.4 算法说明", 3)
    add_para(doc, "TODO: 描述核心算法逻辑。", "宋体", 12, indent=0.74)

    # 示例模块3
    add_heading_styled(doc, "4.3 模块3名称", 2)
    add_heading_styled(doc, "4.3.1 功能说明", 3)
    add_para(doc, "TODO: 描述这个模块的功能。", "宋体", 12, indent=0.74)
    add_heading_styled(doc, "4.3.2 处理流程图", 3)
    _add_image(doc, "06_module3_flow")
    add_heading_styled(doc, "4.3.3 关键函数", 3)
    add_table_from_rows(doc, ["函数名", "所属类", "功能说明"], [
        ("TODO()", "TODO_Service", "TODO功能描述"),
    ], [3, 4, 7])
    add_heading_styled(doc, "4.3.4 算法说明", 3)
    add_para(doc, "TODO: 描述核心算法逻辑。", "宋体", 12, indent=0.74)

    doc.add_page_break()

    # ═════════════════════════════════════════════════════════════════════
    # 五、运行设计
    # ═════════════════════════════════════════════════════════════════════
    add_heading_styled(doc, "五、运行设计", 1)

    add_heading_styled(doc, "5.1 运行流程", 2)
    # TODO: 描述系统的运行流程
    add_para(doc, "TODO: 描述用户使用系统的完整流程。", "宋体", 12, indent=0.74)

    add_heading_styled(doc, "5.2 通信时序", 2)
    add_para(doc, "TODO: 描述系统核心组件间的通信时序。", "宋体", 12, indent=0.74)
    _add_image(doc, "08_sequence")

    add_heading_styled(doc, "5.3 数据流", 2)
    add_para(doc, "TODO: 描述系统的数据流向。", "宋体", 12, indent=0.74)
    _add_image(doc, "10_data_flow")

    doc.add_page_break()

    # ═════════════════════════════════════════════════════════════════════
    # 六、出错处理设计
    # ═════════════════════════════════════════════════════════════════════
    add_heading_styled(doc, "六、出错处理设计", 1)

    add_heading_styled(doc, "6.1 接口异常处理", 2)
    add_para(doc, "TODO: 描述接口异常处理策略。", "宋体", 12, indent=0.74)

    add_heading_styled(doc, "6.2 业务异常处理", 2)
    add_para(doc, "TODO: 描述业务异常处理策略。", "宋体", 12, indent=0.74)

    add_heading_styled(doc, "6.3 系统异常处理", 2)
    add_para(doc, "TODO: 描述系统级异常处理策略。", "宋体", 12, indent=0.74)

    doc.add_page_break()

    # ═════════════════════════════════════════════════════════════════════
    # 七、数据库设计
    # ═════════════════════════════════════════════════════════════════════
    add_heading_styled(doc, "七、数据库设计", 1)

    add_heading_styled(doc, "7.1 ER关系图", 2)
    _add_image(doc, "09_database_er")

    add_heading_styled(doc, "7.2 核心表结构", 2)
    add_para(doc, "TODO: 描述核心数据表。", "宋体", 12, indent=0.74)
    add_table_from_rows(doc, ["表名", "说明", "核心字段"], [
        ("TODO", "TODO说明", "TODO字段"),
    ], [3, 3, 8])

    doc.save(output_path)
    print(f"  [OK] {output_path}")


# ═══════════════════════════════════════════════════════════════════════
# 主入口
# ═══════════════════════════════════════════════════════════════════════

def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    print(f"软件名称: {SOFTWARE_NAME}")
    print(f"版本号: {VERSION}")
    print(f"输出目录: {OUTPUT_DIR}")

    generate_diagrams()
    generate_application_form()
    generate_merged_code()
    generate_design_spec()

    print("\n=== 全部完成 ===")
    print(f"  1. {SOFTWARE_NAME}.doc")
    print(f"  2. {SOFTWARE_NAME}--源代码.docx")
    print(f"  3. {SOFTWARE_NAME}--软件设计说明书.docx")
    print(f"  4. diagrams/ (图表PNG)")


if __name__ == "__main__":
    main()
